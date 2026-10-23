"""Feuille de saisie des absences ScoDoc (« feuille_signal_abs_AAAA-Wnn.xlsx »), générée par le site.

ScoDoc importe une feuille par semaine et par semestre. Format (copié d'une vraie feuille ScoDoc) :
  ligne 1-2  : consignes en rouge ; ligne 3 : titre du semestre (« BUT INFORMATIQUE semestre 1 FI 2026-2027 »)
  ligne 5    : un jour toutes les 2 colonnes à partir de E (« Lundi 19/10/2026 »), samedi et dimanche masqués
  ligne 6    : A = « !id du semestre », puis Nom, Prénom, Groupe, Matin / Après-Midi pour chaque jour
  lignes 7+  : A = « !id ScoDoc de l'étudiant », nom, prénom, groupe, puis ABS / RET dans les cases jaunes

ScoDoc reconnaît les étudiants et le semestre par leurs identifiants internes (les « !… »), que le site
n'a pas au départ. On les relève UNE fois par semestre dans n'importe quelle feuille téléchargée
depuis ScoDoc (lier_feuille) ; ensuite la feuille de chaque semaine se génère directement.

Par demi-journée : ABS si l'étudiant a manqué au moins un cours, sinon RET s'il est arrivé en retard
à au moins un cours, sinon la case reste vide (ScoDoc ne crée rien).
"""
import io
import re
import unicodedata
from collections import Counter
from datetime import datetime, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.worksheet.dimensions import ColumnDimension

import models
from services.fin_annee import PROMOS
from services.presences import ABSENT, RETARD, statuts_etudiants
from services.temps import maintenant_paris

HEURE_MIDI = 12                       # un cours qui commence avant midi compte pour le matin
CODES = {ABSENT: "ABS", RETARD: "RET"}
JOURS = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]
LIGNE_TITRE, LIGNE_JOURS, LIGNE_TITRES = 3, 5, 6
COL_ID, COL_NOM, COL_PRENOM, COL_GROUPE, COL_PREMIER_JOUR = 1, 2, 3, 4, 5


def cle_nom(nom, prenom) -> str:
    """NOM Prénom sans accents, tirets ni majuscules : « Clément » = « Clement »."""
    texte = unicodedata.normalize("NFKD", f"{nom} {prenom}").encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^a-z]+", " ", texte.lower()).split())


def _id(valeur) -> str:
    """« !8725 » -> « 8725 » (vide si la cellule n'est pas un identifiant ScoDoc)."""
    m = re.fullmatch(r"\s*!(\d+)\s*", str(valeur or ""))
    return m[1] if m else ""


def promo_de(etudiant):
    return next((p for p in PROMOS if any(g.nom_groupe.startswith(p) for g in etudiant.groupes)), None)


# ==========================================
# Liaison avec ScoDoc (une fois par semestre)
# ==========================================

def lier_feuille(db, contenu: bytes) -> dict:
    """Relève dans une feuille ScoDoc l'id du semestre et l'id ScoDoc de chaque étudiant
    (reconnu par son nom et son prénom)."""
    ws = load_workbook(io.BytesIO(contenu)).worksheets[0]
    id_semestre = _id(ws.cell(LIGNE_TITRES, COL_ID).value)
    if not id_semestre:
        raise ValueError("Ce fichier ne ressemble pas à une feuille d'absences ScoDoc "
                         "(pas d'identifiant « !… » en case A6).")

    par_nom = {}
    for e in db.query(models.EtudiantDB).all():
        cle = cle_nom(e.nom, e.prenom)
        par_nom[cle] = None if cle in par_nom else e        # homonymes parfaits : pas de liaison

    lies, introuvables, promos = [], [], Counter()
    for r in range(LIGNE_TITRES + 1, ws.max_row + 1):
        nom, prenom, id_etu = ws.cell(r, COL_NOM).value, ws.cell(r, COL_PRENOM).value, _id(ws.cell(r, COL_ID).value)
        if not nom or not id_etu:
            continue
        etudiant = par_nom.get(cle_nom(nom, prenom or ""))
        if etudiant:
            etudiant.id_scodoc = id_etu
            lies.append(etudiant)
            promos[promo_de(etudiant)] += 1
        elif str(ws.cell(r, COL_GROUPE).value or "").strip().upper() != "DEM":    # démissionnaires : normal
            introuvables.append(f"{nom} {prenom or ''}".strip())

    promo = next((p for p, _ in promos.most_common() if p), None)
    if not promo:
        raise ValueError("Aucun étudiant de la feuille n'a été trouvé sur le site : importez d'abord "
                         "les étudiants avec les mêmes noms que dans ScoDoc.")
    titre = str(ws.cell(LIGNE_TITRE, 1).value or promo).strip()
    semestre = db.get(models.ScodocSemestreDB, promo) or models.ScodocSemestreDB(promo=promo)
    semestre.id_semestre, semestre.titre, semestre.modifie_le = id_semestre, titre, maintenant_paris()
    db.merge(semestre)
    db.commit()
    return {"promo": promo, "titre": titre, "lies": len(lies), "introuvables": introuvables}


def etat_liaison(db) -> list:
    """Pour chaque promo : semestre ScoDoc relevé ou non, étudiants liés / total."""
    etudiants = db.query(models.EtudiantDB).all()
    etat = []
    for promo in PROMOS:
        siens = [e for e in etudiants if promo_de(e) == promo]
        semestre = db.get(models.ScodocSemestreDB, promo)
        etat.append({"promo": promo, "titre": semestre.titre if semestre else None,
                     "lies": sum(1 for e in siens if e.id_scodoc), "total": len(siens)})
    return etat


# ==========================================
# Génération de la feuille de la semaine
# ==========================================

def lundi_de(semaine: str = None):
    """« 2026-W43 » (champ semaine du navigateur) -> lundi 00:00 ; par défaut la semaine en cours."""
    if semaine:
        try:
            return datetime.strptime(semaine + "-1", "%G-W%V-%u")
        except ValueError:
            pass
    aujourd_hui = maintenant_paris().replace(hour=0, minute=0, second=0, microsecond=0)
    return aujourd_hui - timedelta(days=aujourd_hui.weekday())


def _groupe(etudiant, promo) -> str:
    """« TD2 / TPD », comme dans ScoDoc."""
    noms = sorted(g.nom_groupe[len(promo) + 1:] for g in etudiant.groupes if g.nom_groupe.startswith(promo + "-"))
    tds = [n for n in noms if n.startswith("TD")]
    tps = [n for n in noms if n.startswith("TP")]
    return " / ".join(tds + tps)


def codes_demi_journees(db, etudiants, debut, fin) -> dict:
    """{(id étudiant, date, après-midi ?): "ABS" | "RET"} ; ABS l'emporte sur RET."""
    codes = {}
    for e, bloc, statut, _, _ in statuts_etudiants(db, etudiants, debut, fin):
        if statut in CODES:
            cle = (e.id, bloc.debut.date(), bloc.debut.hour >= HEURE_MIDI)
            if codes.get(cle) != "ABS":
                codes[cle] = CODES[statut]
    return codes


def generer_feuille(db, promo: str, lundi: datetime, ids_etudiants=None):
    """Feuille ScoDoc de la semaine pour une promo (ou seulement certains de ses étudiants, pour un groupe).
    Renvoie (fichier, nom du fichier, résumé)."""
    tous = [e for e in db.query(models.EtudiantDB).all()
            if promo_de(e) == promo and (ids_etudiants is None or e.id in ids_etudiants)]
    semestre = db.get(models.ScodocSemestreDB, promo)
    if semestre:
        etudiants = [e for e in tous if e.id_scodoc]
    else:
        # Promo non liée : même feuille, avec les numéros du site à la place de ceux de ScoDoc
        # et un titre calculé (semestre impair de septembre à janvier).
        n = int(promo[-1])
        num_semestre = 2 * n - 1 if lundi.month >= 9 or lundi.month == 1 else 2 * n
        annee = lundi.year if lundi.month >= 9 else lundi.year - 1
        semestre = models.ScodocSemestreDB(promo=promo, id_semestre="0",
                                           titre=f"BUT INFORMATIQUE semestre {num_semestre} FI  {annee}-{annee + 1}")
        etudiants = list(tous)
        for e in etudiants:
            e.id_scodoc_affiche = f"0{e.id}"
    etudiants.sort(key=lambda e: (cle_nom(e.nom, ""), cle_nom(e.prenom, "")))
    codes = codes_demi_journees(db, etudiants, lundi, lundi + timedelta(days=7) - timedelta(seconds=1))

    wb = Workbook()
    ws = wb.active
    ws.title = "Saisie_ABS"
    rouge, gras14 = Font(name="Arial", size=12, color="FFFF0000"), Font(name="Arial", size=14, bold=True)
    mauve, normal = Font(name="Arial", size=11, color="FF993366"), Font(name="Arial", size=12)
    case_police, jaune = Font(name="Arial", size=11, bold=True), PatternFill("solid", fgColor="FFFFFF99")
    fin_ = Side(style="thin", color="FF000000")
    bord_case, bord_haut, bord_id = Border(left=fin_, right=fin_, bottom=fin_), Border(top=fin_), Border(right=fin_)

    def ecrire(ligne, col, valeur, police, **style):
        c = ws.cell(ligne, col, valeur)
        c.font = police
        for k, v in style.items():
            setattr(c, k, v)
        return c

    ecrire(1, 1, "Saisir les assiduités dans les cases jaunes (ABS, RET, PRE)", rouge)
    ecrire(2, 1, "Ne pas modifier les cases en mauve !", rouge)
    ecrire(LIGNE_TITRE, 1, semestre.titre, gras14)
    jours = [(lundi + timedelta(days=i)).date() for i in range(7)]
    for col in range(COL_ID, COL_PREMIER_JOUR):
        ecrire(LIGNE_JOURS, col, None, gras14)
    for i, jour in enumerate(jours):
        col = COL_PREMIER_JOUR + 2 * i
        ecrire(LIGNE_JOURS, col, f"{JOURS[i]} {jour:%d/%m/%Y}", gras14)
        ecrire(LIGNE_JOURS, col + 1, None, gras14)
        ecrire(LIGNE_TITRES, col, "Matin", gras14)
        ecrire(LIGNE_TITRES, col + 1, "Après-Midi", gras14)
    ecrire(LIGNE_TITRES, COL_ID, f"!{semestre.id_semestre}", mauve, border=bord_id)
    for col, titre in ((COL_NOM, "Nom"), (COL_PRENOM, "Prénom"), (COL_GROUPE, "Groupe")):
        ecrire(LIGNE_TITRES, col, titre, gras14)

    resume = Counter()
    for r, e in enumerate(etudiants, start=LIGNE_TITRES + 1):
        ecrire(r, COL_ID, f"!{getattr(e, 'id_scodoc_affiche', None) or e.id_scodoc}", mauve, border=bord_id)
        ecrire(r, COL_NOM, e.nom, normal, border=bord_haut)
        ecrire(r, COL_PRENOM, e.prenom, normal, border=bord_haut)
        ecrire(r, COL_GROUPE, _groupe(e, promo), normal, border=bord_haut)
        for i, jour in enumerate(jours):
            for demi in (0, 1):
                code = codes.get((e.id, jour, bool(demi)))
                resume[code] += 1
                ecrire(r, COL_PREMIER_JOUR + 2 * i + demi, code or " ", case_police, fill=jaune, border=bord_case)

    # Mise en page de ScoDoc : largeurs, week-end masqué, hauteurs, filtre
    for lettre, largeur in (("A", 1.5703125), ("B", 23.42578125), ("C", 15.5703125), ("D", 23.42578125)):
        ws.column_dimensions[lettre].width = largeur
    ws.column_dimensions["E"] = ColumnDimension(ws, index="E", min=5, max=14, width=14.28515625)
    ws.column_dimensions["O"] = ColumnDimension(ws, index="O", min=15, max=18, width=14.28515625, hidden=True)
    for ligne in (l for l in range(1, ws.max_row + 1) if l != 4):
        ws.row_dimensions[ligne].height = 18.0 if ligne in (LIGNE_TITRE, LIGNE_JOURS, LIGNE_TITRES) else 15.75
    ws.auto_filter.ref = f"A{LIGNE_TITRES}:R{max(ws.max_row, LIGNE_TITRES)}"

    sortie = io.BytesIO()
    wb.save(sortie)
    sortie.seek(0)
    annee, num, _ = lundi.isocalendar()
    return sortie, f"feuille_signal_abs_{annee}-W{num:02d}.xlsx", {
        "abs": resume["ABS"], "ret": resume["RET"], "etudiants": len(etudiants),
        "sans_id": [f"{e.nom} {e.prenom}" for e in tous if not e.id_scodoc]}
