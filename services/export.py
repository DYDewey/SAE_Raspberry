"""Génération des fichiers Excel (émargement, bilans d'absences)."""
import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def excel_avec_entete(df: pd.DataFrame, feuille: str, entete: list):
    """Tableau Excel précédé de quelques lignes d'en-tête, colonnes ajustées."""
    sortie = io.BytesIO()
    with pd.ExcelWriter(sortie, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=feuille, startrow=len(entete) + 1)
        ws = writer.sheets[feuille]
        for i, ligne in enumerate(entete, start=1):
            ws.cell(row=i, column=1, value=ligne).font = Font(bold=(i == 1), size=14 if i == 1 else 11)
        for cellule in ws[len(entete) + 2]:
            cellule.font = Font(bold=True, color="FFFFFF")
            cellule.fill = PatternFill("solid", fgColor="343A40")
        for col in ws.iter_cols(min_row=len(entete) + 2):
            largeur = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = max(12, largeur + 2)
    sortie.seek(0)
    return sortie


JOURS = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]
COULEUR_STATUT = {"Absent": "F8D7DA", "En retard": "FFF3CD"}


def excel_absences(lignes, entete: list):
    """Une ligne par absence ou retard, prête à reporter dans ScoDoc :
    filtres sur chaque colonne, en-tête figé, statut coloré."""
    df = pd.DataFrame([{
        "Date": l["bloc"].debut.strftime("%d/%m/%Y"),
        "Jour": JOURS[l["bloc"].debut.weekday()],
        "Début": l["bloc"].debut.strftime("%H:%M"),
        "Fin": l["bloc"].fin.strftime("%H:%M"),
        "Durée (h)": round((l["bloc"].fin - l["bloc"].debut).total_seconds() / 3600, 2),
        "Code cours": l["bloc"].code,
        "Cours": l["bloc"].libelle,
        "Groupe": l["bloc"].nom_groupe,
        "Salle": l["bloc"].nom_salle,
        "Professeur": l["bloc"].nom_prof,
        "N° étudiant": l["etudiant"].numero_etudiant,
        "Nom": l["etudiant"].nom,
        "Prénom": l["etudiant"].prenom,
        "Statut": l["statut"],
        "Arrivée": l["arrivee"].strftime("%H:%M") if l["arrivee"] else "",
        "Retard (min)": int((l["arrivee"] - l["bloc"].debut).total_seconds() // 60) if l["arrivee"] else "",
        "Saisie": "manuelle" if l["manuel"] else "badge",
    } for l in lignes], columns=["Date", "Jour", "Début", "Fin", "Durée (h)", "Code cours", "Cours", "Groupe",
                                 "Salle", "Professeur", "N° étudiant", "Nom", "Prénom", "Statut", "Arrivée",
                                 "Retard (min)", "Saisie"])
    sortie = excel_avec_entete(df, "Absences", entete)

    # Finitions : filtres, en-tête figé, statut coloré
    from openpyxl import load_workbook
    wb = load_workbook(sortie)
    ws = wb["Absences"]
    ligne_titres = len(entete) + 2
    ws.auto_filter.ref = f"A{ligne_titres}:{ws.cell(row=ligne_titres, column=ws.max_column).column_letter}{ws.max_row}"
    ws.freeze_panes = ws.cell(row=ligne_titres + 1, column=1)
    col_statut = df.columns.get_loc("Statut") + 1
    for (cellule,) in ws.iter_rows(min_row=ligne_titres + 1, min_col=col_statut, max_col=col_statut):
        if cellule.value in COULEUR_STATUT:
            cellule.fill = PatternFill("solid", fgColor=COULEUR_STATUT[cellule.value])
            cellule.font = Font(bold=True)
            cellule.alignment = Alignment(horizontal="center")
    sortie = io.BytesIO()
    wb.save(sortie)
    sortie.seek(0)
    return sortie
